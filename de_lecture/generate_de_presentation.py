import os
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = pptx.Presentation()
    # 16:9 widescreen layout
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6] # completely blank layout

    # Color Palette
    BG_COLOR = RGBColor(11, 15, 25)        # Midnight Navy #0B0F19
    CARD_BG = RGBColor(15, 23, 42)         # Deep Slate #0F172A
    CYAN = RGBColor(56, 189, 248)          # #38BDF8
    AMBER = RGBColor(245, 158, 11)         # #F59E0B
    GREEN = RGBColor(16, 185, 129)         # #10B981
    CORAL = RGBColor(244, 63, 94)          # #F43F5E
    TEXT_WHITE = RGBColor(248, 250, 252)   # #F8FAFC
    MUTED_GRAY = RGBColor(148, 163, 184)   # #94A3B8

    SLIDES_DATA = [
        {
            "num": "01 / 21",
            "title": "Applications of Differential Equations",
            "subtitle": "Elementary Applications of First-Order ODEs // Mathematical Modeling in Motion",
            "frame": "public/frames/Scene01_IntroDE.png",
            "video": "public/videos/Scene01_IntroDE.mp4",
            "key_takeaways": [
                "Dynamic Differential Modeling vs Static Snapshot Algebra",
                "Rates of change define future physical trajectories",
                "Primary source: MATHalino Engineering Mathematics Review"
            ],
            "notes": "Introduce the lecture scope. Contrast static functions y = f(x) with dynamic differential equations dy/dt = f(t, y). Explain how rates of change define physical laws."
        },
        {
            "num": "02 / 21",
            "title": "The 4-Step Mathematical Modeling Blueprint",
            "subtitle": "From Empirical Observation to Analytical Prediction",
            "frame": "public/frames/Scene02_ModelingBlueprint.png",
            "video": "public/videos/Scene02_ModelingBlueprint.mp4",
            "key_takeaways": [
                "Step 1: State the physical proportionality law",
                "Step 2: Formulate the differential equation with constant k",
                "Step 3: Separate variables and integrate analytically",
                "Step 4: Apply boundary/initial conditions to evaluate C"
            ],
            "notes": "Walk through the universal 4-step framework. Emphasize that all rate-driven physical processes (thermal, nuclear, mechanical) follow this identical workflow."
        },
        {
            "num": "03 / 21",
            "title": "Separation of Variables Mathematical Engine",
            "subtitle": "Transforming Differential Statements into Integrated Solution Families",
            "frame": "public/frames/Scene03_SeparationVariables.png",
            "video": "public/videos/Scene03_SeparationVariables.mp4",
            "key_takeaways": [
                "Standard separable form: dy/dt = g(y) h(t)",
                "Group dependent terms on left: dy / g(y) = h(t) dt",
                "Integrate both sides: ∫ dy/g(y) = ∫ h(t) dt + C",
                "Solve explicitly for y(t) using inversion/logarithms"
            ],
            "notes": "Demonstrate the algebraic mechanics of separation. Highlight how Leibniz differentials dy and dt are separated across the equality sign."
        },
        {
            "num": "04 / 21",
            "title": "Newton's Law of Cooling: Physical Intuition",
            "subtitle": "Thermal Driving Potential & Asymptotic Decay Toward Ambient Temperature",
            "frame": "de_lecture/cooling_final_frame.png",
            "video": "media/videos/prototype_cooling/720p30/NewtonCoolingPrototype.mp4",
            "key_takeaways": [
                "Governing law: dT/dt = -k(T - Ts)",
                "Driving force is the temperature difference (T - Ts)",
                "Rate is fastest initially and decreases asymptotically",
                "Ambient temperature Ts serves as a horizontal asymptote"
            ],
            "notes": "Explain why dT/dt is proportional to (T - Ts). When T > Ts, the body cools (dT/dt < 0). Point out that the curve never crosses Ts under pure convective cooling."
        },
        {
            "num": "05 / 21",
            "title": "Newton's Law of Cooling: Analytical Derivation",
            "subtitle": "Continuous Equation Morphing from Rate Law to Explicit Solution",
            "frame": "public/frames/Scene04_NewtonCoolingLaw.png",
            "video": "public/videos/Scene04_NewtonCoolingLaw.mp4",
            "key_takeaways": [
                "Separated form: dT / (T - Ts) = -k dt",
                "Integrating: ln(T - Ts) = -kt + ln C1",
                "Exponentiating: T - Ts = C1 e^(-kt)",
                "Initial condition: T(0) = T0 ⟹ C1 = T0 - Ts",
                "Particular model: T(t) = Ts + (T0 - Ts) e^(-kt)"
            ],
            "notes": "Step through each line of the derivation. Emphasize how the initial condition evaluates the constant C1 as the initial temperature excess."
        },
        {
            "num": "06 / 21",
            "title": "Mathalino Problem 01: Thermometer Reading After 6 Minutes",
            "subtitle": "Given: Inside T₀ = 70°F, Outside Ts = 10°F, T(3) = 25°F. Find T(6).",
            "frame": "public/frames/Scene05_CoolingProblem1_Thermometer6Min.png",
            "video": "public/videos/Scene05_CoolingProblem1_Thermometer6Min.mp4",
            "key_takeaways": [
                "Model: T(t) = 10 + 60 e^(-kt)",
                "At t = 3: 25 = 10 + 60 e^(-3k) ⟹ e^(-3k) = 1/4",
                "Base model: T(t) = 10 + 60 (1/4)^(t/3)",
                "Evaluation at t = 6: T(6) = 10 + 60(1/16) = 13.75°F"
            ],
            "notes": "State MATHalino Problem 01 data. Ask students to deduce e^(-3k) = 1/4 before moving to the resolution. Highlight the constant ratio drop over equal 3-minute intervals."
        },
        {
            "num": "07 / 21",
            "title": "Problem 01 Resolution & Ratio Invariance",
            "subtitle": "Constant Interval Scaling Principle in First-Order Linear Systems",
            "frame": "public/frames/Scene05_CoolingProblem1_Thermometer6Min.png",
            "video": None,
            "key_takeaways": [
                "Interval 0 to 3 min: ΔT drops from 60°F to 15°F (ratio = 1/4)",
                "Interval 3 to 6 min: ΔT drops from 15°F to 3.75°F (ratio = 1/4)",
                "Final Temperature: T(6) = 10 + 3.75 = 13.75°F exactly",
                "Verified against MATHalino official engineering solution"
            ],
            "notes": "Use Apple Pencil on iPad to annotate the ratio arrows 60 -> 15 -> 3.75 on the temperature axis."
        },
        {
            "num": "08 / 21",
            "title": "Mathalino Problem 02: Time to Within Half Degree of Air",
            "subtitle": "T₀ = 75°F, Ts = 20°F, T(4) = 30°F. Target: T = 20.5°F. Find elapsed time t.",
            "frame": "public/frames/Scene06_CoolingProblem2_HalfDegree.png",
            "video": "public/videos/Scene06_CoolingProblem2_HalfDegree.mp4",
            "key_takeaways": [
                "Model: T(t) = 20 + 55 e^(-kt)",
                "At t = 4: 30 = 20 + 55 e^(-4k) ⟹ e^(-4k) = 10/55 = 2/11",
                "Target condition: T(t) - 20 = 0.5°F ⟹ T = 20.5°F",
                "Equation: 0.5 = 55 (2/11)^(t/4) ⟹ 1/110 = (2/11)^(t/4)"
            ],
            "notes": "Explain why the target is 20.5°F (within half a degree of 20°F). Emphasize that reaching 20°F exactly would take infinite time."
        },
        {
            "num": "09 / 21",
            "title": "Problem 02 Logarithmic Resolution",
            "subtitle": "Isolating Unknown Time Exponent via Natural Logarithms",
            "frame": "public/frames/Scene06_CoolingProblem2_HalfDegree.png",
            "video": None,
            "key_takeaways": [
                "Logarithmic form: ln(1/110) = (t/4) ln(2/11)",
                "Exact solution: t = 4 ln(1/110) / ln(2/11)",
                "Numerical value: t ≈ 11.029 minutes",
                "Independently verified against MATHalino records"
            ],
            "notes": "Show how ln(1/110) = -ln 110 and ln(2/11) = ln 2 - ln 11 are both negative, giving a positive elapsed time of 11.029 minutes."
        },
        {
            "num": "10 / 21",
            "title": "Thermal Modeling Live Derivation Board",
            "subtitle": "Instructor Apple Pencil Inking Space // Parameter Variations",
            "frame": None,
            "video": None,
            "key_takeaways": [
                "Explore heating scenarios where T0 < Ts (dT/dt > 0)",
                "Examine how convection coefficients alter cooling rate k",
                "Practice problem: Predict reading at t = 9 minutes"
            ],
            "notes": "Use this slide as a live whiteboard. Pose a question to students: What if the object was brought inside from the cold into a warm room?"
        },
        {
            "num": "11 / 21",
            "title": "Simple Chemical Conversion & Radioactive Decay",
            "subtitle": "First-Order Dissipation Kinetics & Unconverted Substance Remaining",
            "frame": "public/frames/Scene07_ChemicalConversion_Decay.png",
            "video": "public/videos/Scene07_ChemicalConversion_Decay.mp4",
            "key_takeaways": [
                "Governing law: dx/dt = -kx (k > 0)",
                "x(t) represents remaining unconverted/undecayed substance",
                "Separation: dx / x = -k dt ⟹ ln x = -kt + C",
                "General solution: x(t) = x0 e^(-kt)"
            ],
            "notes": "Introduce chemical conversion and nuclear decay. Clarify that the negative sign indicates loss of reactant or parent isotope."
        },
        {
            "num": "12 / 21",
            "title": "Stochastic Atoms vs Continuum Calculus",
            "subtitle": "Bridging Quantum Indeterminacy and Macroscopic Exponential Decay",
            "frame": "de_lecture/decay_final_frame.png",
            "video": "media/videos/prototype_decay/720p30/RadioactiveDecayPrototype.mp4",
            "key_takeaways": [
                "Individual decay events are quantum and probabilistic",
                "Avogadro-scale ensemble (10²³ atoms) yields smooth deterministic ODE",
                "Remaining fraction x(t)/x0 follows strict exponential decay",
                "Half-life t1/2 is independent of initial sample mass x0"
            ],
            "notes": "Discuss why we can use a smooth continuous differential equation for discrete atomic decay. Explain the Law of Large Numbers."
        },
        {
            "num": "13 / 21",
            "title": "Mathalino Decay Problem 01: Radium Half-Life",
            "subtitle": "Given: 1.1% Decomposes in 25 Years (x(25) = 0.989 x₀). Find Half-Life.",
            "frame": "public/frames/Scene08_RadiumProblem1_HalfLife.png",
            "video": "public/videos/Scene08_RadiumProblem1_HalfLife.mp4",
            "key_takeaways": [
                "Decomposed fraction: 1.1% ⟹ Remaining fraction: 98.9% (0.989)",
                "Condition at t = 25: 0.989 x0 = x0 e^(-25k) ⟹ e^(-k) = (0.989)^(1/25)",
                "Decay model: x(t) = x0 (0.989)^(t/25)",
                "Half-life definition: x(t) = 0.50 x0"
            ],
            "notes": "Highlight MATHalino Radium Problem 01. Warn students not to confuse 1.1% decomposed with 1.1% remaining."
        },
        {
            "num": "14 / 21",
            "title": "Radium Half-Life Century-Scale Resolution",
            "subtitle": "Calibrating a Multi-Millennium Geological Clock from a 25-Year Sample",
            "frame": "public/frames/Scene08_RadiumProblem1_HalfLife.png",
            "video": None,
            "key_takeaways": [
                "Setting x = 0.5 x0: 0.5 = (0.989)^(t/25)",
                "Taking logarithms: 25 ln(0.5) = t ln(0.989)",
                "Analytical solution: t = 25 ln(0.5) / ln(0.989)",
                "Calculated half-life: t ≈ 1,566.65 years",
                "Matches MATHalino reference solution exactly"
            ],
            "notes": "Demonstrate the calculation. Explain that Radium-226 has a known half-life of ~1600 years, showing the high physical accuracy of MATHalino's data."
        },
        {
            "num": "15 / 21",
            "title": "Mathalino Decay Problem 02: 90% Radioactivity Dissipation",
            "subtitle": "Given: Half-Life t₁/₂ = 38 Hours. Find Time for 90% to Dissipate.",
            "frame": "public/frames/Scene09_DecayProblem2_90Percent.png",
            "video": "public/videos/Scene09_DecayProblem2_90Percent.mp4",
            "key_takeaways": [
                "Base model using half-life: x(t) = x0 (0.5)^(t/38)",
                "Critical terminology: 90% dissipated ⟹ 10% remaining (x = 0.10 x0)",
                "Equation: 0.10 x0 = x0 (0.5)^(t/38) ⟹ 0.10 = 0.5^(t/38)",
                "Target time: t = 38 ln(0.10) / ln(0.5) ≈ 126.23 hours"
            ],
            "notes": "Present MATHalino Problem 02. Emphasize: 90% dissipated means exactly 10% remaining! Ask students to estimate how many half-lives this requires."
        },
        {
            "num": "16 / 21",
            "title": "Problem 02 Threshold Resolution & Milestone Check",
            "subtitle": "Verification Against Consecutive Half-Life Milestones",
            "frame": "public/frames/Scene09_DecayProblem2_90Percent.png",
            "video": None,
            "key_takeaways": [
                "t = 38 hr (1 half-life): 50% remaining",
                "t = 76 hr (2 half-lives): 25% remaining",
                "t = 114 hr (3 half-lives): 12.5% remaining",
                "t ≈ 126.23 hr (3.32 half-lives): 10.0% remaining",
                "Verified against MATHalino official engineering solution"
            ],
            "notes": "Annotate the half-life milestones on the slide. Showing that 126.23 hr lies between 3 and 4 half-lives gives students immediate physical confidence in the result."
        },
        {
            "num": "17 / 21",
            "title": "Engineering Physics Extension: Velocity of Escape from Earth",
            "subtitle": "Primary MATHalino Index Topic // Vertical Motion Under Variable Gravity",
            "frame": "public/frames/Scene10_EscapeVelocity.png",
            "video": "public/videos/Scene10_EscapeVelocity.mp4",
            "key_takeaways": [
                "Newton's Gravitation: F = -m g R² / r²",
                "Kinematic Chain Rule: a = dv/dt = v (dv/dr)",
                "Governing DE: v (dv/dr) = -g R² / r²",
                "Separation of variables: v dv = -g R² r^(-2) dr"
            ],
            "notes": "Transition to the engineering extension topic from MATHalino. Explain why acceleration a is written as v dv/dr rather than dv/dt."
        },
        {
            "num": "18 / 21",
            "title": "Escape Velocity Analytical Resolution",
            "subtitle": "Integrating to Infinity to Overcome the Planetary Gravitational Well",
            "frame": "public/frames/Scene10_EscapeVelocity.png",
            "video": None,
            "key_takeaways": [
                "Integrating: ∫ v dv = -g R² ∫ r^(-2) dr ⟹ (1/2)v² = g R² / r + C",
                "Boundary condition at infinity: r ➔ ∞, v ≥ 0 ⟹ C ≥ 0",
                "Surface launch (r = R): (1/2) v_esc² = g R ⟹ v_esc = √(2gR)",
                "Numerical value: v_esc ≈ 11.18 km/s (6.95 mi/s ≈ 25,000 mph)"
            ],
            "notes": "Derive v_esc = sqrt(2gR). Show how energy conservation and differential separation yield the identical escape formula."
        },
        {
            "num": "19 / 21",
            "title": "Master Comparative Modeling Matrix",
            "subtitle": "Unified Differential Architecture Across Engineering Applications",
            "frame": "public/frames/Scene11_SynthesisComparison.png",
            "video": "public/videos/Scene11_SynthesisComparison.mp4",
            "key_takeaways": [
                "Cooling: dT/dt = -k(T - Ts) ⟹ Asymptotic thermal balance",
                "Decay: dx/dt = -kx ⟹ Invariant half-life dissipation",
                "Escape: v dv/dr = -gR²/r² ⟹ Gravitational potential escape",
                "All 3 solve cleanly via separation of variables"
            ],
            "notes": "Summarize the 3 models side-by-side. Highlight the unity of first-order differential equations across physical systems."
        },
        {
            "num": "20 / 21",
            "title": "Lecture Summary & Live Apple Pencil Q&A",
            "subtitle": "Core Engineering Takeaways & Examination Review",
            "frame": None,
            "video": None,
            "key_takeaways": [
                "Always check units and physical domain constraints",
                "Distinguish remaining fraction from dissipated fraction",
                "Identify physical asymptotes before solving transcendental equations"
            ],
            "notes": "Open the floor for questions. Provide live derivations for any requested steps."
        },
        {
            "num": "21 / 21",
            "title": "Curriculum References & Academic Acknowledgements",
            "subtitle": "Primary Sources, Literature References & Animation Framework",
            "frame": None,
            "video": None,
            "key_takeaways": [
                "MATHalino Engineering Math Review (mathalino.com)",
                "James Stewart, Calculus: Early Transcendentals (9th Ed.)",
                "Boyce & DiPrima, Elementary Differential Equations",
                "3Blue1Brown (Grant Sanderson), Essence of Calculus style"
            ],
            "notes": "Credit MATHalino for all authentic source problems and numerical data."
        }
    ]

    for s_idx, data in enumerate(SLIDES_DATA):
        slide = prs.slides.add_slide(blank_layout)

        # Background Fill
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background()

        # Top Banner Text
        tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(1.1))
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p_num = tf.paragraphs[0]
        p_num.text = f"SLIDE {data['num']} // DIFFERENTIAL EQUATIONS"
        p_num.font.size = Pt(11)
        p_num.font.bold = True
        p_num.font.color.rgb = CYAN

        p_title = tf.add_paragraph()
        p_title.text = data["title"]
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_WHITE
        p_title.space_before = Pt(3)

        p_sub = tf.add_paragraph()
        p_sub.text = data["subtitle"]
        p_sub.font.size = Pt(13)
        p_sub.font.color.rgb = MUTED_GRAY
        p_sub.space_before = Pt(2)

        # Layout Split: Frame/Video on Left, Derivation / Annotations on Right
        has_media = data["frame"] and os.path.exists(data["frame"])

        if has_media:
            # Place Still Frame on Left
            slide.shapes.add_picture(data["frame"], Inches(0.8), Inches(1.8), width=Inches(6.4))

            # Right Card: Key Takeaways & Apple Pencil Annotation Zone
            right_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.5), Inches(1.8), Inches(5.0), Inches(5.1))
            right_card.fill.solid()
            right_card.fill.fore_color.rgb = CARD_BG
            right_card.line.color.rgb = CYAN
            right_card.line.width = Pt(1.5)

            rt_box = slide.shapes.add_textbox(Inches(7.7), Inches(2.0), Inches(4.6), Inches(4.7))
            rtf = rt_box.text_frame
            rtf.word_wrap = True

            p_h = rtf.paragraphs[0]
            p_h.text = "KEY MATHEMATICAL TAKEAWAYS"
            p_h.font.size = Pt(14)
            p_h.font.bold = True
            p_h.font.color.rgb = AMBER

            for item in data["key_takeaways"]:
                p_it = rtf.add_paragraph()
                p_it.text = f"• {item}"
                p_it.font.size = Pt(12)
                p_it.font.color.rgb = TEXT_WHITE
                p_it.space_before = Pt(8)

            p_ann = rtf.add_paragraph()
            p_ann.text = " Apple Pencil Annotation Zone"
            p_ann.font.size = Pt(12)
            p_ann.font.bold = True
            p_ann.font.color.rgb = GREEN
            p_ann.space_before = Pt(16)

            p_ann_sub = rtf.add_paragraph()
            p_ann_sub.text = "(Generous writing area for live derivations and class questions)"
            p_ann_sub.font.size = Pt(10.5)
            p_ann_sub.font.color.rgb = MUTED_GRAY

        else:
            # Full-Width Card for derivations / summary / references
            full_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.1))
            full_card.fill.solid()
            full_card.fill.fore_color.rgb = CARD_BG
            full_card.line.color.rgb = CYAN if s_idx != 20 else AMBER
            full_card.line.width = Pt(1.5)

            ft_box = slide.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(11.1), Inches(4.5))
            ftf = ft_box.text_frame
            ftf.word_wrap = True

            p_h = ftf.paragraphs[0]
            p_h.text = "LECTURE FOCUS & CORE TAKEAWAYS"
            p_h.font.size = Pt(16)
            p_h.font.bold = True
            p_h.font.color.rgb = AMBER

            for item in data["key_takeaways"]:
                p_it = ftf.add_paragraph()
                p_it.text = f"• {item}"
                p_it.font.size = Pt(14)
                p_it.font.color.rgb = TEXT_WHITE
                p_it.space_before = Pt(12)

        # Speaker Notes directly attached to the slide
        notes_slide = slide.notes_slide
        text_frame = notes_slide.notes_text_frame
        text_frame.text = f"TEACHER NOTES & GUIDED QUESTIONS (Slide {data['num']}):\n{data['notes']}"

    output_dir = "presentation"
    os.makedirs(output_dir, exist_ok=True)
    out_path = os.path.join(output_dir, "Differential_Equations_Lecture.pptx")
    prs.save(out_path)
    print(f"Successfully generated PowerPoint presentation deck: {out_path}")

    # Also save to public for immediate web serving / download
    public_dir = "public"
    os.makedirs(public_dir, exist_ok=True)
    public_path = os.path.join(public_dir, "Differential_Equations_Lecture.pptx")
    prs.save(public_path)
    print(f"Copied PowerPoint presentation deck to: {public_path}")

if __name__ == "__main__":
    create_deck()
